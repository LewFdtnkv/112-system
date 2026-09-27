from datetime import UTC, datetime, timedelta
from types import SimpleNamespace
from uuid import UUID, uuid4

import pytest
from fastapi import HTTPException
from sqlalchemy import select
from test_catalog_dds import api as api
from test_catalog_dds import dds as dds
from test_catalog_dds import document
from test_dds_crews import crews as crews
from test_teacher_api import teaching as teaching

from app.models import Evaluation
from app.schemas.learning import LearningPolicy
from app.services.dds.timing import timing
from app.services.group_reports import aggregate
from app.services.learning_scope import validate_exercise

pytestmark = pytest.mark.anyio


@pytest.mark.parametrize(
    "opened,record,states",
    [
        (30, 180, ("on_time", "on_time")),
        (31, 181, ("late", "late")),
        (None, None, ("missing", "missing")),
    ],
)
async def test_separate_clocks_and_pauses(opened, record, states):
    start = datetime(2026, 1, 1, tzinfo=UTC)
    attempt = SimpleNamespace(
        settings_snapshot={"dds_policy": {"workflow": "crews-v2"}},
        started_at=start,
        ended_at=start + timedelta(seconds=500),
        pauses=[],
        first_opened_at=start + timedelta(seconds=opened) if opened is not None else None,
        first_record_at=start + timedelta(seconds=record) if record is not None else None,
    )
    result = timing(attempt)
    assert tuple(v["state"] for v in result.values()) == states
    attempt.pauses = [
        {
            "start": (start + timedelta(seconds=10)).isoformat(),
            "end": (start + timedelta(seconds=20)).isoformat(),
        }
    ]
    if record:
        assert timing(attempt)["first_record"]["seconds"] == record - 10
    attempt.settings_snapshot["dds_policy"]["workflow"] = "crews-v1"
    assert timing(attempt) is None


@pytest.fixture
async def work(crews, api):
    c = crews
    t = c.d.t
    exercise = {
        "workflow": "crews-v2",
        "service_profile_id": c.profile["id"],
        "crew_calls_required": False,
        "initial_crews": [],
        "required_crews": [{"crew_code": "water", "status": "completed"}],
        "messages": [
            {
                "crew_code": "water",
                "message": "Бригада приняла, выехала, прибыла, "
                "приступила к устранению утечки и завершила работы.",
            }
        ],
    }
    card = await t.post(
        "cards",
        t.card_payload
        | {"recipient_service_ids": c.d.card["recipient_service_ids"], "dds_exercise": exercise},
    )
    scenario = await t.post(
        "scenarios",
        {
            "title": "QA4",
            "role": "dds",
            "service_profile_id": c.profile["id"],
            "card_ids": [card["id"]],
        },
    )
    lesson = await t.post(
        "lessons/start",
        {
            "request_id": str(uuid4()),
            "group_id": c.d.group["id"],
            "scenario_version_id": scenario["id"],
            "learning": {"kind": "practice"},
        },
    )
    journal = await api("GET", f"student/lessons/{lesson['id']}", actor="student")
    a = await api(
        "POST",
        f"student/assignments/{journal['assignments'][0]['id']}/start",
        actor="student",
        status=201,
    )
    path = f"student/attempts/{a['id']}"

    async def act(a, state, comment="Подтверждено руководителем бригады", status=200):
        return await api(
            "POST",
            path + "/dds/crews",
            c.command(a, status=state, comment=comment),
            actor="student",
            status=status,
        )

    return SimpleNamespace(**locals())


async def test_v2_full_path_clocks_and_unchanged_service(work, api):
    w = work
    a = w.a
    assert a["dds"]["workflow"] == "crews-v2"
    assert a["dds"]["timing"]["opening"]["at"]
    service_times = [(r["status"], r["status_updated_at"]) for r in a["dds"]["responses"]]
    a = await w.act(a, "assigned", "")
    assert a["dds"]["timing"]["first_record"]["at"] is None
    await w.act(a, "responding", status=422)
    await w.act(a, "accepted", "   ", status=422)
    a = await w.act(a, "accepted")
    first = a["dds"]["timing"]["first_record"]["at"]
    assert first
    for status in ("responding", "arrived", "in_progress", "completed"):
        if status == "in_progress":
            await w.act(a, "completed", status=422)
        a = await w.act(a, status)
        assert a["dds"]["timing"]["first_record"]["at"] == first
    assert [(r["status"], r["status_updated_at"]) for r in a["dds"]["responses"]] == service_times
    await api("POST", w.path + "/dds/submit", {"revision": a["dds"]["revision"]}, actor="student")
    grade = await w.t.db_session.scalar(
        select(Evaluation).where(Evaluation.attempt_id == UUID(a["id"]))
    )
    assert grade.score == grade.max_score
    assert grade.context_snapshot["dds"]["timing"]["first_record"]["state"] == "on_time"


async def test_refusal_has_reason_and_is_not_cancellation(work):
    w = work
    a = await w.act(w.a, "assigned")
    await w.act(a, "not_accepted", "", status=422)
    a = await w.act(a, "not_accepted", "Нет необходимой техники")
    assert a["dds"]["crews"][0]["status"] == "not_accepted"
    a = await w.act(a, "assigned")
    a = await w.act(a, "accepted")
    await w.act(a, "not_accepted", "Причина", status=422)
    await w.act(a, "refused", " ", status=422)
    a = await w.act(a, "refused", "Обнаружена опасность для бригады")
    assert a["dds"]["crews"][0]["status"] == "refused"


async def test_full_situation_rejects_intermediate_goal_but_skill_accepts():
    exercise = {
        "workflow": "crews-v2",
        "initial_crews": [],
        "required_crews": [{"crew_code": "water", "status": "assigned"}],
    }
    scenario = SimpleNamespace(role="dds", completion_rules={})
    cards = [SimpleNamespace(snapshot={"dds_exercise": exercise})]
    with pytest.raises(HTTPException):
        validate_exercise(LearningPolicy(kind="practice"), scenario, cards)
    validate_exercise(
        LearningPolicy(kind="skill_practice", target_skills=["dds_crews"]), scenario, cards
    )


async def test_territory_roundtrip_and_teacher_preview(api):
    doc = document()
    doc["entries"][0]["routes"][1]["addresses"] = [
        {"locality": "Москва", "street": "Тверская"},
        {"locality": "Москва", "district": "Арбат"},
    ]
    version = await api("POST", "admin/classifiers/import", doc, status=201)
    path = f"admin/classifiers/{version['id']}"
    entry = (await api("GET", path + "/entries"))["items"][0]["id"]
    await api("POST", path + "/publish")
    exported = await api("GET", path + "/export")
    assert (
        exported["entries"][0]["routes"][1]["addresses"]
        == doc["entries"][0]["routes"][1]["addresses"]
    )
    preview = f"classifiers/{version['id']}/recipients-preview"
    body = {
        "classifier_entry_id": entry,
        "answers": {"victims": True},
        "address": {"locality": "МОСКВА", "street": "тверская"},
    }
    matched = await api("POST", preview, body, actor="teacher")
    assert len(matched) == 2
    for address in ({}, {"locality": "Тверь", "street": "Тверская"}):
        assert len(await api("POST", preview, body | {"address": address}, actor="teacher")) == 1
    assert (
        len(await api("POST", preview, body | {"answers": {"victims": False}}, actor="teacher"))
        == 1
    )
    await api("POST", preview, body, actor="student", status=403)


async def test_group_aggregate_deduplicates_cards_and_respects_teacher():
    base = {
        "status": "completed",
        "teacher_review": None,
        "kind": "practice",
        "student_id": "s1",
        "lesson_id": "l1",
        "assignment_id": "a1",
        "position": 1,
        "credits": {"dds_response": 0.5},
    }
    cards = [
        base,
        base | {"assignment_id": "a2"},
        base | {"assignment_id": "a3", "student_id": "s2", "credits": {"dds_response": 1}},
        base | {"teacher_review": "Верно", "assignment_id": "a4"},
    ]
    result = aggregate(cards)
    stat = result["statistics"][0]
    assert (stat["students"], stat["affected"], stat["cards"], stat["failed_cards"]) == (2, 1, 3, 2)
    assert result["teacher_reviewed_cards"] == 1
    assert result["profile"]["candidates"] == ["dds_response"]


async def test_report_group_queue_permissions_and_fenced_publication(work, api, db_client):
    from io import BytesIO

    from openpyxl import load_workbook

    from app.models import AIJob
    from app.services import group_reports
    from app.services.generation_worker import finish
    from app.services.learning_recommendations.materials import retrieve

    w = work
    # Finish without doing the work: exact missing-action evidence, not an LLM failure.
    await api("POST", w.path + "/dds/submit", {"revision": w.a["dds"]["revision"]}, actor="student")
    path = f"teaching/groups/{w.c.d.group['id']}/analysis?role=dds&days=30"
    report = await api("GET", path, actor="teacher")
    assert report["statistics"] and any(s["failed_cards"] for s in report["statistics"])
    await api("GET", path, actor="student", status=403)
    await api("GET", path, actor="other", status=404)
    queued = await api("POST", path, actor="teacher", status=202)
    assert await api("POST", path, actor="teacher", status=202) == queued
    job = await w.t.db_session.get(AIJob, UUID(queued["id"]))
    token = str(uuid4())
    job.status = "running"
    job.worker_id = token
    job.lease_expires_at = datetime.now(UTC) + timedelta(minutes=5)
    bundle = await retrieve(w.t.db_session, job.input["profile"], vectors=False)
    job.context = {"materials": bundle}
    await w.t.db_session.commit()
    ids = {}
    for example in bundle["examples"]:
        ids.setdefault(example["skill"], example["id"])
    output = {"selected_ids": list(ids.values()), "mode": "ai"}
    job_id = job.id
    assert not await finish(w.t.db_session, job_id, str(uuid4()), output, {})
    assert await finish(w.t.db_session, job_id, token, output, {})
    report = await api("GET", path, actor="teacher")
    assert report["job"]["status"] == "succeeded"
    assert report["job"]["recommendations"] and not report["job"]["obsolete"]
    response = await db_client.get(
        "/api/v1/teaching/reports/export", headers=w.t.headers["teacher"]
    )
    assert response.status_code == 200, response.text
    workbook = load_workbook(BytesIO(response.content))
    assert {"Карточки", "Ошибки и комментарии"} <= set(workbook.sheetnames)
    rows = list(workbook["Карточки"].values)
    assert "Норматив записи, с" in rows[0]
    assert any("Не выполнено" in row for row in rows[1:])
    assert group_reports.PROMPT_VERSION == job.prompt_version


async def test_semantic_evidence_not_inferred_from_failures():
    from app.services.learning_recommendations.profile import credits_for

    field = {"field": "description", "scored": True, "status": "matched"}
    criteria = [SimpleNamespace(criterion_snapshot={"fields": [field]})]
    job = SimpleNamespace(
        input={"criteria": []},
        output={"findings": [{"code": "description", "applied": True, "credit": 0.5}]},
    )
    assert credits_for(None, criteria, job) == {"description": 0.5}
    job.output = {"findings": [{"code": "description", "applied": False, "credit": 0}]}
    assert credits_for(None, criteria, job) == {}
    job.output = {}
    assert credits_for(None, criteria, job) == {}
