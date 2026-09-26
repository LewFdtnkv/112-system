from datetime import UTC, datetime, timedelta
from io import BytesIO
from uuid import UUID, uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import func, select
from test_student_workflow import exercise as exercise
from test_teacher_api import teaching as teaching

from app.models import (
    Attempt,
    AttemptEvent,
    Lesson,
    LessonEvaluation,
    ProctoringEvent,
    Scenario,
    ServiceResponse,
    UserActivity,
)
from app.services.deadlines import enforce_deadlines, sweep_deadlines

pytestmark = pytest.mark.anyio


async def test_deadline_scores_partial_and_missing_without_notification(exercise, db_session):
    e = exercise
    filled = await e.fill(await e.start())
    lesson = await db_session.get(Lesson, UUID(e.lesson["id"]))
    deadline = datetime.now(UTC) + timedelta(seconds=1)
    lesson.available_until = deadline
    await db_session.commit()
    await enforce_deadlines(db_session, lesson, deadline + timedelta(seconds=1))
    attempt = await db_session.get(Attempt, UUID(filled["id"]))
    assert attempt.status == "interrupted" and attempt.ended_at == deadline
    assert not await db_session.scalar(select(func.count()).select_from(ServiceResponse))
    grade = await db_session.scalar(
        select(LessonEvaluation).where(LessonEvaluation.lesson_id == lesson.id)
    )
    assert grade and 0 < grade.score < 34
    assert grade.assessment_details["missed_cards"] == 2
    exposed = await e.request("GET", f"student/lessons/{lesson.id}/evaluation")
    assert exposed["assessment_details"]["missed_cards"] == 2
    assert (await e.request("GET", f"student/lessons/{lesson.id}"))["work_status"] == "submitted"
    await e.request(
        "PUT",
        f"student/attempts/{attempt.id}/card",
        {"revision": filled["card"]["revision"], "data": {}},
        status=409,
    )
    await sweep_deadlines(db_session)
    assert await db_session.scalar(select(func.count()).select_from(LessonEvaluation)) == 1
    response = await e.request(
        "POST",
        f"lessons/{lesson.id}/students/{attempt.student_id}/evaluations",
        {
            "request_id": str(uuid4()),
            "expected_revision": 1,
            "score": 50,
            "max_score": 100,
            "comment": "Пересмотр",
        },
        actor="teacher",
        status=201,
    )
    assert response["method"] == "teacher"


async def test_future_start_and_entirely_missed_assignment(teaching, db_client, db_session):
    t = teaching
    d = await t.prepare(members=("student",))
    start = datetime.now(UTC) + timedelta(hours=1)
    end = start + timedelta(minutes=10)
    lesson = await t.post(
        "lessons/start",
        d.payload | {"available_from": start.isoformat(), "available_until": end.isoformat()},
    )
    assert lesson["status"] == "planned"
    work = (
        await db_client.get(f"/api/v1/student/lessons/{lesson['id']}", headers=t.headers["student"])
    ).json()
    response = await db_client.post(
        f"/api/v1/student/assignments/{work['assignments'][0]['id']}/start",
        headers=t.headers["student"],
    )
    assert response.status_code == 409
    stored = await db_session.get(Lesson, UUID(lesson["id"]))
    await enforce_deadlines(db_session, stored, end + timedelta(seconds=1))
    result = await db_session.scalar(
        select(LessonEvaluation).where(LessonEvaluation.lesson_id == stored.id)
    )
    assert result.score == 0 and result.assessment_details["missed_cards"] == 3
    assert not await db_session.scalar(select(func.count()).select_from(Attempt))


async def test_card_timeout_unlocks_next_and_keeps_lesson_open(exercise, db_session):
    e = exercise
    first = await e.start()
    row = await db_session.get(Attempt, UUID(first["id"]))
    row.started_at = datetime.now(UTC) - timedelta(minutes=4)
    await db_session.commit()
    work = await e.request("GET", f"student/lessons/{e.lesson['id']}")
    assert work["status"] == "active"
    assert work["assignments"][0]["status"] == "interrupted"
    assert work["assignments"][1]["available"]


async def test_proctoring_is_separate_scoped_and_idempotent(exercise, db_session):
    e = exercise
    attempt = await e.start()
    event = {
        "command_id": str(uuid4()),
        "kind": "tab.hidden",
        "client_occurred_at": datetime.now(UTC).isoformat(),
    }
    path = f"student/attempts/{attempt['id']}/proctoring"
    await e.request("POST", path, {"events": [event]})
    await e.request("POST", path, {"events": [event]})
    await e.request("POST", path, {"events": [event | {"kind": "tab.visible"}]}, status=409)
    await e.request("POST", path, {"events": [event]}, actor="student2", status=404)
    await e.request(
        "GET", f"teaching/attempts/{attempt['id']}/proctoring", actor="other", status=404
    )
    history = await e.request(
        "GET", f"teaching/attempts/{attempt['id']}/proctoring", actor="teacher"
    )
    assert history["total"] == 1 and history["trusted"] is False
    monitor = await e.request("GET", "teaching/monitoring", actor="teacher")
    assert monitor["items"][0]["student_id"] == str(e.t.accounts["student"].id)
    assert monitor["items"][0]["visibility"] == "tab.hidden"
    assert monitor["items"][0]["hidden_count"] == 1
    assert await db_session.scalar(select(func.count()).select_from(ProctoringEvent)) == 1
    assert not await db_session.scalar(
        select(func.count()).select_from(AttemptEvent).where(AttemptEvent.kind.like("tab.%"))
    )
    await e.request(
        "POST", f"student/attempts/{attempt['id']}/observations", {"events": [event]}, status=422
    )


async def test_group_messages_survive_transfer_and_reports_are_scoped(teaching, db_client):
    t = teaching
    d = await t.prepare(members=("student",))
    await t.post("lessons/start", d.payload)
    message = await t.post("messages", {"group_id": d.group["id"], "text": "=Учебное объявление"})
    target = await t.post("groups", {"name": "Новая группа"})
    path = f"/api/v1/groups/{d.group['id']}/students/{t.accounts['student'].id}/transfer"
    assert (
        await db_client.post(
            path, headers=t.headers["other"], json={"target_group_id": target["id"]}
        )
    ).status_code == 404
    assert (
        await db_client.post(
            path, headers=t.headers["teacher"], json={"target_group_id": target["id"]}
        )
    ).status_code == 204
    inbox = (await db_client.get("/api/v1/student/messages", headers=t.headers["student"])).json()
    assert inbox["total"] == 1 and inbox["items"][0]["id"] == message["id"]
    assert (await db_client.get("/api/v1/student/messages", headers=t.headers["student2"])).json()[
        "total"
    ] == 0
    report = await db_client.get(
        f"/api/v1/teaching/reports/export?student_id={t.accounts['student'].id}",
        headers=t.headers["teacher"],
    )
    assert report.status_code == 200
    assert load_workbook(BytesIO(report.content)).active.max_row == 2
    assert (
        await db_client.get(
            f"/api/v1/teaching/students/{t.accounts['student'].id}", headers=t.headers["other"]
        )
    ).status_code == 404


async def test_account_reason_and_export(teaching, db_client, db_session):
    t = teaching
    user_id = t.accounts["student"].id
    path = f"/api/v1/users/{user_id}"
    assert (
        await db_client.patch(path, headers=t.headers["admin"], json={"is_active": False})
    ).status_code == 422
    assert (
        await db_client.patch(
            path, headers=t.headers["admin"], json={"is_active": False, "reason": "=Причина"}
        )
    ).status_code == 200
    event = await db_session.scalar(
        select(UserActivity).where(
            UserActivity.user_id == user_id, UserActivity.kind == "account.access_changed"
        )
    )
    assert event.reason == "=Причина" and event.details["after"]["is_active"] is False
    report = await db_client.get(
        f"/api/v1/admin/users/{user_id}/activity/export?format=xlsx", headers=t.headers["admin"]
    )
    assert report.status_code == 200
    cells = list(load_workbook(BytesIO(report.content)).active.values)
    assert any("=Причина" in row for row in cells)
    assert (
        await db_client.get(f"/api/v1/admin/users/{user_id}/activity", headers=t.headers["teacher"])
    ).status_code == 403


async def test_multi_target_dedup_and_scenario_archive(teaching, db_client, db_session):
    t = teaching
    d = await t.prepare()
    payload = {k: v for k, v in d.payload.items() if k != "group_id"}
    lesson = await t.post(
        "lessons/start",
        payload | {"group_ids": [d.group["id"]], "student_ids": [str(t.accounts["student"].id)]},
    )
    assert lesson["student_count"] == 2 and lesson["assignment_count"] == 6
    response = await db_client.delete(
        f"/api/v1/scenarios/{d.scenario['id']}", headers=t.headers["teacher"]
    )
    assert response.json() == {"result": "archived"}
    parent = await db_session.get(Scenario, UUID(d.scenario["scenario_id"]))
    assert parent.is_archived
    await t.post("lessons/start", d.payload | {"request_id": str(uuid4())}, expected=409)
    assert (
        await db_client.get(f"/api/v1/lessons/{lesson['id']}", headers=t.headers["teacher"])
    ).status_code == 200
    other = await t.prepare()
    response = await db_client.delete(
        f"/api/v1/scenarios/{other.scenario['id']}", headers=t.headers["teacher"]
    )
    assert response.json() == {"result": "deleted"}


async def test_admin_statistics_and_photo_permissions(teaching, db_client):
    from PIL import Image

    t = teaching
    stats = await db_client.get("/api/v1/admin/statistics", headers=t.headers["admin"])
    assert stats.status_code == 200
    assert sum(row["registered"] for row in stats.json()) >= 5
    user_id = t.accounts["student"].id
    path = f"/api/v1/admin/users/{user_id}/photo"
    buffer = BytesIO()
    Image.new("RGB", (40, 40), color="red").save(buffer, "PNG")
    response = await db_client.put(path, headers=t.headers["admin"], content=buffer.getvalue())
    assert response.status_code == 204
    photo = await db_client.get(f"/api/v1/users/{user_id}/photo", headers=t.headers["student"])
    assert photo.status_code == 200 and photo.headers["content-type"] == "image/jpeg"
    assert (
        await db_client.get(f"/api/v1/users/{user_id}/photo", headers=t.headers["student2"])
    ).status_code == 404
    assert (
        await db_client.put(path, headers=t.headers["admin"], content=b"<svg/>")
    ).status_code == 422
